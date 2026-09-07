import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_element(name):
    return OxmlElement(name)

def create_attribute(element, name, value):
    element.set(qn(name), value)

def add_page_number(run):
    fldChar1 = create_element('w:fldChar')
    create_attribute(fldChar1, 'w:fldCharType', 'begin')
    instrText = create_element('w:instrText')
    create_attribute(instrText, 'xml:space', 'preserve')
    instrText.text = "PAGE"
    fldChar2 = create_element('w:fldChar')
    create_attribute(fldChar2, 'w:fldCharType', 'separate')
    fldChar3 = create_element('w:fldChar')
    create_attribute(fldChar3, 'w:fldCharType', 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = docx.Document()

    # Set Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set Default Font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Add Footer Page Numbers
    footer = doc.sections[0].footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = footer_p.add_run("Multi-Tier Cloud Application Project Report | Page ")
    f_run.font.size = Pt(9)
    f_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    add_page_number(f_run)

    lines = content.split('\n')
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "\n".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        run = cp.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        code_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        # Parse table markdown
        rows_data = []
        for tl in table_lines:
            if re.match(r'^\s*\|?\s*:?-+:?\s*\|', tl):
                continue # delimiter row
            cols = [c.strip() for c in tl.strip('|').split('|')]
            if any(cols):
                rows_data.append(cols)

        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=len(rows_data), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for r_idx, row in enumerate(rows_data):
                for c_idx, val in enumerate(row):
                    if c_idx < num_cols:
                        cell = table.cell(r_idx, c_idx)
                        cell.text = val
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(4)
                        
                        if r_idx == 0:
                            set_cell_background(cell, "1E293B")
                            for run in p.runs:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F8FAFC")
                            else:
                                set_cell_background(cell, "FFFFFF")

        table_lines = []

    i = 0
    while i < len(lines):
        line = lines[i]

        # Handle Code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                flush_code_block()
                in_code_block = False
            else:
                if in_table:
                    flush_table()
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if '|' in line and line.strip().startswith('|'):
            if not in_table:
                in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()
            in_table = False

        # Page Break
        if '<div page-break-after="always"></div>' in line or line.strip() == '***':
            doc.add_page_break()
            i += 1
            continue

        # Align Center divs
        if '<div align="center">' in line:
            # Parse title block until </div>
            i += 1
            while i < len(lines) and '</div>' not in lines[i]:
                tline = lines[i].strip()
                if tline:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    if tline.startswith('# '):
                        run = p.add_run(tline.replace('# ', ''))
                        run.font.size = Pt(22)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                    elif tline.startswith('### '):
                        run = p.add_run(tline.replace('### ', ''))
                        run.font.size = Pt(13)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                    elif tline.startswith('## '):
                        run = p.add_run(tline.replace('## ', ''))
                        run.font.size = Pt(16)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
                    else:
                        cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', tline)
                        run = p.add_run(cleaned)
                        run.font.size = Pt(11)
                        if '**' in tline:
                            run.font.bold = True
                i += 1
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('# ', ''))
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('## ', ''))
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('### ', ''))
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:]
            # Format bold inside bullet
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)

        i += 1

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Successfully generated Word Document: {docx_path}")

if __name__ == "__main__":
    md_file = r"C:\Users\ganag\.gemini\antigravity\scratch\multi-tier-cloud-ecommerce\PROJECT_REPORT.md"
    docx_file = r"C:\Users\ganag\.gemini\antigravity\scratch\multi-tier-cloud-ecommerce\PROJECT_REPORT.docx"
    convert_md_to_docx(md_file, docx_file)
